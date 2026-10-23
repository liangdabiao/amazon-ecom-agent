"""广告报表工具：解析 SP 搜索词报表（xlsx/csv）→ 聚合统计 → 交给 agent 分析。

原 IvyeaOps 的做法是把整份报表塞进提示词让外部 agent 跑；这里反过来：
工具负责确定性的读表和算数（这个不能出错），诊断结论由对话 agent 依据
系统提示词里的方法论产出。
"""
from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

from agents import function_tool

from ..config import config

# 报表列名 → 规范字段（亚马搜索词报表的中英文表头都覆盖）
_COLUMN_ALIASES: dict[str, list[str]] = {
    "campaign": ["campaign", "campaign name", "广告活动", "广告活动名称", "推广活动"],
    "ad_group": ["ad group", "ad group name", "广告组", "广告组名称"],
    "term": ["customer search term", "search term", "search terms", "搜索词", "客户搜索词", "关键词"],
    "match_type": ["match type", "匹配类型"],
    "impressions": ["impressions", "展示量", "曝光量"],
    "clicks": ["clicks", "点击量", "点击次数"],
    "ctr": ["ctr", "点击率"],
    "cpc": ["cpc", "cost per click", "平均点击费用"],
    "spend": ["spend", "cost", "花费", "总花费"],
    "orders": ["7 day total orders", "orders", "total orders", "订单量", "7天总订单数", "订单数"],
    "sales": ["7 day total sales", "sales", "total sales", "销售额", "7天总销售额"],
    "acos": ["acos", "广告投入产出比"],
}


def _num(v: Any) -> float:
    if v is None or isinstance(v, bool):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").replace("$", "").replace("%", "").strip() or 0)
    except ValueError:
        return 0.0


def _match_columns(headers: list[str]) -> dict[str, str]:
    """表头模糊匹配到规范字段。"""
    norm = {h: str(h or "").strip().lower() for h in headers}
    out: dict[str, str] = {}
    for field, aliases in _COLUMN_ALIASES.items():
        for alias in aliases:
            for h, low in norm.items():
                if low == alias or (alias in low and len(alias) >= 4):
                    out.setdefault(field, h)
                    break
            if field in out:
                break
    return out


def _read_rows(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    """读 xlsx / csv，返回 (headers, rows)。xlsx 用 read_only 模式防大表爆内存。"""
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(str(path), read_only=True, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        headers = [str(c or "") for c in next(rows_iter)]
        rows = []
        for r in rows_iter:
            if r is None or all(c is None or str(c).strip() == "" for c in r):
                continue
            rows.append({headers[i]: ("" if i >= len(r) or r[i] is None else r[i])
                         for i in range(len(headers))})
        wb.close()
        return headers, rows
    if suffix == ".csv":
        text = path.read_text("utf-8-sig", errors="replace")
        sample = text[:4096]
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;") if sample else csv.excel
        reader = csv.DictReader(text.splitlines(), dialect=dialect)
        headers = reader.fieldnames or []
        return headers, [dict(r) for r in reader]
    raise RuntimeError(f"不支持的报表格式: {path.suffix}（支持 xlsx / csv）")


def _aggregate(headers: list[str], rows: list[dict[str, str]], col: dict[str, str]) -> dict:
    total = {"impressions": 0.0, "clicks": 0.0, "spend": 0.0, "orders": 0.0, "sales": 0.0}
    campaigns: dict[str, dict] = {}
    terms: list[dict] = []
    c = lambda r, f: _num(r.get(col[f], 0)) if f in col else 0.0  # noqa: E731

    for r in rows:
        rec = {f: c(r, f) for f in total}
        for k in total:
            total[k] += rec[k]
        camp = str(r.get(col["campaign"], "(未知活动)")) if "campaign" in col else "(未知活动)"
        slot = campaigns.setdefault(camp, {**{k: 0.0 for k in total}, "terms": 0})
        for k in total:
            slot[k] += rec[k]
        slot["terms"] += 1
        if "term" in col:
            terms.append({
                "campaign": camp,
                "term": str(r.get(col["term"], ""))[:100],
                "match": str(r.get(col.get("match_type", ""), "")),
                "impressions": int(rec["impressions"]),
                "clicks": int(rec["clicks"]),
                "spend": round(rec["spend"], 2),
                "orders": int(rec["orders"]),
                "sales": round(rec["sales"], 2),
            })

    def _ratio(a: float, b: float) -> str:
        return f"{a / b * 100:.1f}%" if b > 0 else ("N/A" if a > 0 else "0.0%")

    camp_stats = []
    for name, s in sorted(campaigns.items(), key=lambda kv: -kv[1]["spend"]):
        camp_stats.append({
            "campaign": name,
            "spend": round(s["spend"], 2),
            "clicks": int(s["clicks"]),
            "orders": int(s["orders"]),
            "sales": round(s["sales"], 2),
            "acos": _ratio(s["spend"], s["sales"]),
            "cvr": _ratio(s["orders"], s["clicks"]),
            "cost_per_order": round(s["spend"] / s["orders"], 2) if s["orders"] else None,
        })

    terms.sort(key=lambda t: -t["spend"])
    # 花了钱没出单的词放最前（这是诊断最需要的），其余按花费排
    wasted = [t for t in terms if t["spend"] > 0 and t["orders"] == 0]
    performing = [t for t in terms if t["orders"] > 0]
    return {
        "total": {
            **{k: round(v, 2) for k, v in total.items()},
            "acos": _ratio(total["spend"], total["sales"]),
            "ctr": _ratio(total["clicks"], total["impressions"]),
            "cvr": _ratio(total["orders"], total["clicks"]),
            "term_rows": len(rows),
        },
        "campaigns": camp_stats,
        "top_wasted_terms": wasted[:30],
        "top_performing_terms": performing[:30],
    }


@function_tool
def list_report_files() -> str:
    """列出 data/reports/ 目录下的广告报表文件（xlsx/csv）。
    用户上传报表后就放这里；先用本工具看有哪些文件，再用 analyze_ads_report 分析。"""
    files = []
    for p in sorted(config.reports_dir.iterdir()):
        if p.suffix.lower() in (".xlsx", ".xlsm", ".csv") and p.is_file():
            files.append({"file": p.name, "path": str(p),
                          "size_kb": round(p.stat().st_size / 1024, 1)})
    if not files:
        return "data/reports/ 目录为空。请把亚马逊广告搜索词报表（xlsx/csv）放进去再试。"
    return json.dumps(files, ensure_ascii=False, indent=2)


@function_tool
def analyze_ads_report(file_name: str) -> str:
    """解析一份广告搜索词报表（xlsx/csv），返回确定性聚合：总消耗/ACOS/CTR/CVR、
    每个 Campaign 的效率对比、烧钱无单的搜索词 Top30、出单搜索词 Top30。
    拿到本工具返回的数据后，按系统提示词里的广告诊断方法论出结论。

    Args:
        file_name: 报表文件名（list_report_files 列出的名字，如 report.xlsx）
    """
    path = config.reports_dir / file_name
    if not path.is_file():
        path = config.reports_dir / (file_name if file_name.endswith((".xlsx", ".csv"))
                                     else file_name + ".xlsx")
    if not path.is_file():
        return f"找不到报表文件 {file_name}，请先用 list_report_files 查看可用文件。"
    try:
        headers, rows = _read_rows(path)
        col = _match_columns(headers)
        missing = [f for f in ("campaign", "spend", "clicks") if f not in col]
        if missing:
            return (f"报表 {file_name} 缺少关键列（{', '.join(missing)}）。\n"
                    f"识别到的表头：{headers[:20]}")
        result = _aggregate(headers, rows, col)
        result["file"] = file_name
        result["matched_columns"] = col
        return json.dumps(result, ensure_ascii=False, indent=2)
    except Exception as exc:
        return f"解析报表 {file_name} 失败: {exc}"
